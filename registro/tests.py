from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .forms import EstudianteForm
from .models import Estudiante
from .validators import validar_cedula_ecuatoriana


DATOS_VALIDOS = {"cedula": "1710034065", "nombre": "Ana María", "apellido": "Pérez Mora", "correo": "ana@example.com", "edad": 20, "carrera": Estudiante.Carrera.SOFTWARE}


def crear_archivo_excel(filas, encabezados=None):
    libro = Workbook()
    hoja = libro.active
    hoja.append(encabezados or ["cedula", "nombre", "apellido", "correo", "edad", "carrera"])
    for fila in filas:
        hoja.append(fila)
    contenido = BytesIO()
    libro.save(contenido)
    libro.close()
    return SimpleUploadedFile(
        "estudiantes.xlsx",
        contenido.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class ValidadorCedulaTests(TestCase):
    def test_acepta_cedula_ecuatoriana_valida(self):
        validar_cedula_ecuatoriana("1710034065")

    def test_rechaza_cedula_invalida(self):
        with self.assertRaises(ValidationError):
            validar_cedula_ecuatoriana("0912345678")


class EstudianteFormTests(TestCase):
    def test_formulario_valido_normaliza_datos(self):
        datos = {**DATOS_VALIDOS, "nombre": "  ana   maría ", "correo": "ANA@EXAMPLE.COM"}
        form = EstudianteForm(datos)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["nombre"], "Ana María")
        self.assertEqual(form.cleaned_data["correo"], "ana@example.com")

    def test_rechaza_edad_fuera_del_rango(self):
        for edad in (15, 101):
            with self.subTest(edad=edad):
                form = EstudianteForm({**DATOS_VALIDOS, "edad": edad})
                self.assertFalse(form.is_valid())
                self.assertIn("edad", form.errors)


class EstudianteViewsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser(username="admin", password="clave-segura", email="admin@example.com")
        cls.usuario_sin_permisos = get_user_model().objects.create_user(username="consulta", password="clave-segura")
        cls.estudiante = Estudiante.objects.create(**DATOS_VALIDOS)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_usuario_anonimo_es_redirigido_al_login(self):
        self.client.logout()
        respuesta = self.client.get(reverse("registro:lista"))
        self.assertRedirects(respuesta, f"{reverse('login')}?next={reverse('registro:lista')}")

    def test_usuario_sin_permiso_recibe_403(self):
        self.client.force_login(self.usuario_sin_permisos)
        self.assertEqual(self.client.get(reverse("registro:lista")).status_code, 403)

    def test_administracion_usa_identidad_visual_de_la_app(self):
        respuesta = self.client.get(reverse("admin:index"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "Gestión académica")
        self.assertContains(respuesta, "Administración")
        self.assertContains(respuesta, "css/admin.css")
        self.assertContains(respuesta, "Volver al sistema")

    def test_lista_y_busqueda(self):
        respuesta = self.client.get(reverse("registro:lista"), {"q": "Pérez"})
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "Ana María")
        self.assertEqual(respuesta.context["pagina"].paginator.count, 1)

    def test_registra_estudiante_valido(self):
        datos = {**DATOS_VALIDOS, "cedula": "0926687856", "correo": "nuevo@example.com"}
        respuesta = self.client.post(reverse("registro:registrar"), datos)
        creado = Estudiante.objects.get(cedula="0926687856")
        self.assertRedirects(respuesta, reverse("registro:detalle", args=[creado.pk]))

    def test_edicion_invalida_no_modifica_estudiante(self):
        respuesta = self.client.post(reverse("registro:editar", args=[self.estudiante.pk]), {"correo": "ana@example.com", "edad": 101, "carrera": "ADM"})
        self.assertEqual(respuesta.status_code, 200)
        self.estudiante.refresh_from_db()
        self.assertEqual(self.estudiante.edad, 20)

    def test_edicion_valida(self):
        respuesta = self.client.post(reverse("registro:editar", args=[self.estudiante.pk]), {"correo": "actualizado@example.com", "edad": 21, "carrera": "ADM"})
        self.assertRedirects(respuesta, reverse("registro:detalle", args=[self.estudiante.pk]))
        self.estudiante.refresh_from_db()
        self.assertEqual(self.estudiante.correo, "actualizado@example.com")

    def test_eliminacion_exige_post(self):
        url = reverse("registro:eliminar", args=[self.estudiante.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url).status_code, 302)
        self.assertFalse(Estudiante.objects.filter(pk=self.estudiante.pk).exists())

    def test_importa_excel_valido(self):
        archivo = crear_archivo_excel(
            [["0926687856", "Luis", "Mora", "luis@example.com", 22, "ADM"]]
        )
        respuesta = self.client.post(
            reverse("registro:importar_excel"), {"archivo": archivo}
        )
        self.assertRedirects(respuesta, reverse("registro:lista"))
        importado = Estudiante.objects.get(cedula="0926687856")
        self.assertEqual(importado.nombre, "Luis")
        self.assertEqual(importado.correo, "luis@example.com")

    def test_rechaza_todo_el_excel_si_una_fila_es_invalida(self):
        archivo = crear_archivo_excel(
            [
                ["0926687856", "Luis", "Mora", "luis@example.com", 22, "ADM"],
                ["0912345678", "Invalido", "Mora", "malo@example.com", 15, "SOF"],
            ]
        )
        respuesta = self.client.post(
            reverse("registro:importar_excel"), {"archivo": archivo}
        )
        self.assertEqual(respuesta.status_code, 400)
        self.assertContains(respuesta, "Fila 3", status_code=400)
        self.assertFalse(Estudiante.objects.filter(cedula="0926687856").exists())
        self.assertFalse(Estudiante.objects.filter(correo="malo@example.com").exists())

    def test_rechaza_duplicados_en_el_excel(self):
        archivo = crear_archivo_excel(
            [
                ["0926687856", "Luis", "Mora", "luis@example.com", 22, "ADM"],
                ["0926687856", "Luisa", "Paz", "luisa@example.com", 24, "DER"],
            ]
        )
        respuesta = self.client.post(
            reverse("registro:importar_excel"), {"archivo": archivo}
        )
        self.assertEqual(respuesta.status_code, 400)
        self.assertContains(respuesta, "repetida en el archivo", status_code=400)
        self.assertEqual(Estudiante.objects.count(), 1)

    def test_exporta_excel_reutilizable(self):
        respuesta = self.client.get(reverse("registro:exportar_excel"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        libro = load_workbook(BytesIO(respuesta.content), read_only=True)
        filas = list(libro.active.iter_rows(values_only=True))
        libro.close()
        self.assertEqual(
            filas[0], ("cedula", "nombre", "apellido", "correo", "edad", "carrera")
        )
        self.assertEqual(filas[1][0], self.estudiante.cedula)
