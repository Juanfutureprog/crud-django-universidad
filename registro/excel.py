from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .forms import EstudianteForm
from .models import Estudiante


ENCABEZADOS = ("cedula", "nombre", "apellido", "correo", "edad", "carrera")
ETIQUETAS = ENCABEZADOS


class ErrorImportacionExcel(ValidationError):
    pass


def _texto(valor):
    return "" if valor is None else str(valor).strip()


def _cedula(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip().zfill(10)


def importar_estudiantes(archivo):
    try:
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorImportacionExcel("El archivo no es un Excel .xlsx válido.") from exc

    try:
        hoja = libro.active
        filas = hoja.iter_rows(values_only=True)
        encabezados = next(filas, None)
        encabezados_normalizados = tuple(_texto(valor).lower() for valor in (encabezados or ()))
        if encabezados_normalizados != ENCABEZADOS:
            raise ErrorImportacionExcel(
                "Los encabezados deben ser, en este orden: " + ", ".join(ENCABEZADOS) + "."
            )

        formularios = []
        errores = []
        cedulas = set()
        correos = set()
        for numero_fila, fila in enumerate(filas, start=2):
            if all(valor is None or _texto(valor) == "" for valor in fila):
                continue
            valores = list(fila[: len(ENCABEZADOS)])
            valores.extend([None] * (len(ENCABEZADOS) - len(valores)))
            datos = {
                "cedula": _cedula(valores[0]),
                "nombre": _texto(valores[1]),
                "apellido": _texto(valores[2]),
                "correo": _texto(valores[3]).lower(),
                "edad": valores[4],
                "carrera": _texto(valores[5]).upper(),
            }
            if datos["cedula"] in cedulas:
                errores.append(f"Fila {numero_fila}: la cédula está repetida en el archivo.")
            if datos["correo"] in correos:
                errores.append(f"Fila {numero_fila}: el correo está repetido en el archivo.")
            cedulas.add(datos["cedula"])
            correos.add(datos["correo"])

            formulario = EstudianteForm(datos)
            if not formulario.is_valid():
                detalle = "; ".join(
                    f"{campo}: {', '.join(mensajes)}"
                    for campo, mensajes in formulario.errors.items()
                )
                errores.append(f"Fila {numero_fila}: {detalle}")
            formularios.append(formulario)

        if not formularios:
            errores.append("El archivo no contiene estudiantes.")
        if errores:
            raise ErrorImportacionExcel(errores)

        with transaction.atomic():
            Estudiante.objects.bulk_create(
                [formulario.save(commit=False) for formulario in formularios]
            )
        return len(formularios)
    finally:
        libro.close()


def crear_excel_estudiantes(estudiantes):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Estudiantes"
    hoja.append(ETIQUETAS)
    for estudiante in estudiantes:
        hoja.append(
            (
                estudiante.cedula,
                estudiante.nombre,
                estudiante.apellido,
                estudiante.correo,
                estudiante.edad,
                estudiante.carrera,
            )
        )

    color = "8F1720"
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(horizontal="center")
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    anchos = (15, 25, 25, 32, 10, 14)
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.column_dimensions["A"].number_format = "@"

    salida = BytesIO()
    libro.save(salida)
    libro.close()
    return salida.getvalue()
